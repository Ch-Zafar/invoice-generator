import re
from pathlib import Path
import openpyxl


def get_excel_file_path(filename: str = "Sales Purchase Register.xlsx") -> Path:
    """Locates the Excel file in the current working directory or project root."""
    path = Path(filename)
    if path.is_file():
        return path

    project_root = Path(__file__).resolve().parent.parent.parent
    root_path = project_root / filename
    if root_path.is_file():
        return root_path

    raise FileNotFoundError(f"Excel file '{filename}' was not found.")


def get_next_invoice_number(filename: str = "Sales Purchase Register.xlsx") -> int:
    """Opens the 'Sale Register' sheet and returns the next invoice number (last invoice # + 1)."""
    excel_path = get_excel_file_path(filename)

    # Load workbook with data_only=True to read evaluated values instead of formulas
    workbook = openpyxl.load_workbook(excel_path, data_only=True)
    try:
        # 1. Find the Sales sheet (matches 'Sale Register' or 'Sale')
        sheet = None
        for name in workbook.sheetnames:
            if "sale" in name.lower():
                sheet = workbook[name]
                break

        if sheet is None:
            raise ValueError(f"Could not find a sales sheet in {workbook.sheetnames}")

        # 2. Locate the 'Inv. No.' column and header row
        invoice_col = None
        header_row = None

        for row in range(1, 15):
            for col in range(1, sheet.max_column + 1):
                val = sheet.cell(row=row, column=col).value
                if val and any(keyword in str(val).lower() for keyword in ["inv. no", "invoice #", "inv #", "invoice no"]):
                    invoice_col = col
                    header_row = row
                    break
            if invoice_col:
                break

        # Fallback to Column F (col index 6) and Row 4 if header search misses
        if not invoice_col:
            invoice_col = 6
            header_row = 4

        # 3. Read the table rows to find the last valid invoice number
        last_invoice_str = None
        for row in range(header_row + 1, sheet.max_row + 1):
            cell_value = sheet.cell(row=row, column=invoice_col).value
            if cell_value is not None:
                cleaned = str(cell_value).strip()
                if cleaned:
                    last_invoice_str = cleaned

        if not last_invoice_str:
            raise ValueError(f"No invoice numbers found in column {invoice_col} of sheet '{sheet.title}'.")

        # 4. Extract the integer number and return last invoice # + 1
        match = re.search(r"\d+", last_invoice_str)
        if not match:
            raise ValueError(f"Could not parse numeric digits from invoice number: '{last_invoice_str}'")

        last_number = int(match.group())
        next_number = last_number + 1

        return next_number

    finally:
        workbook.close()


if __name__ == "__main__":
    next_inv = get_next_invoice_number()
    print(f"Next Invoice Number: {next_inv}")
